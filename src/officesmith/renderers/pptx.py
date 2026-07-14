from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from ..schemas import PresentationSpec

# Layout indexes in python-pptx's default template.
_TITLE_LAYOUT = 0
_BULLETS_LAYOUT = 1
_SECTION_LAYOUT = 2


def render_presentation(spec: PresentationSpec, out_dir: Path) -> Path:
    from . import output_path

    prs = Presentation()

    for slide_spec in spec.slides:
        if slide_spec.layout == "title":
            layout = prs.slide_layouts[_TITLE_LAYOUT]
        elif slide_spec.layout == "section":
            layout = prs.slide_layouts[_SECTION_LAYOUT]
        else:
            layout = prs.slide_layouts[_BULLETS_LAYOUT]

        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = slide_spec.title

        if slide_spec.layout in ("title", "section") and slide_spec.subtitle:
            slide.placeholders[1].text = slide_spec.subtitle

        if slide_spec.layout == "bullets" and slide_spec.bullets:
            body = slide.placeholders[1].text_frame
            for i, bullet in enumerate(slide_spec.bullets):
                para = body.paragraphs[0] if i == 0 else body.add_paragraph()
                para.text = bullet.text
                para.level = bullet.level

        if slide_spec.notes:
            slide.notes_slide.notes_text_frame.text = slide_spec.notes

    path = output_path(out_dir, spec.filename, ".pptx")
    prs.save(str(path))
    return path
